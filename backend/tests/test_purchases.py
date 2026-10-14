"""Subscriptions & Purchases — money math, cycle stepping and the xlsx report.
DB-free: the pieces that decide numbers people put in front of finance."""
import io
from datetime import date, timedelta
from types import SimpleNamespace

import pytest
from openpyxl import load_workbook
from pydantic import ValidationError

from app.services import purchases_service as ps


@pytest.fixture
def fx(monkeypatch):
    def set_table(raw):
        monkeypatch.setattr(ps, "eff", lambda key, default, cast=str: raw if key == "PURCHASES_FX_TO_AED" else default)
    return set_table


def test_usd_is_the_peg_and_cannot_be_overridden(fx):
    fx("USD=1,EUR=4")
    assert ps.to_aed(100, "USD") == 367.25
    assert ps.aed_to_usd(367.25) == 100.0
    assert ps.to_aed(10, "eur") == 40.0


def test_unknown_or_malformed_rate_is_none_not_a_guess(fx):
    fx("EUR=abc,RUB=0.044,XX=3")
    assert ps.to_aed(100, "EUR") is None
    assert ps.to_aed(100, "RUB") == 4.4
    assert ps.to_aed(100, "JPY") is None


@pytest.mark.parametrize("start,months,expected", [
    (date(2026, 1, 31), 1, date(2026, 2, 28)),
    (date(2028, 1, 31), 1, date(2028, 2, 29)),
    (date(2026, 11, 15), 3, date(2027, 2, 15)),
    (date(2026, 10, 15), 12, date(2027, 10, 15)),
    (date(2026, 3, 1), -11, date(2025, 4, 1)),
])
def test_add_months_clamps_to_month_end(start, months, expected):
    assert ps.add_months(start, months) == expected


def _sub(**kw):
    base = dict(id=1, vendor="Adobe", product="CC", category="software", seller=None, owner_user_ids=None,
                budget=None, budget_id=None, account_email=None, payment_method="card", card=None, card_id=None,
                cycle_months=1, price=100.0, currency="AED", seats_total=None, seats=[], next_renewal=None,
                auto_renew=True, status="active", notify_days_before=None, reference=None, notes=None)
    base.update(kw)
    return SimpleNamespace(**base)


def test_monthly_run_rate_spreads_yearly_price(fx):
    fx("")
    assert ps.monthly_aed(_sub(price=1200, cycle_months=12)) == 100.0
    assert ps.monthly_aed(_sub(price=100, currency="USD")) == 367.25


def test_lapsed_only_when_not_auto_renewing():
    past = date.today() - timedelta(days=1)
    assert ps.is_lapsed(_sub(auto_renew=False, next_renewal=past))
    assert not ps.is_lapsed(_sub(auto_renew=True, next_renewal=past))
    assert not ps.is_lapsed(_sub(auto_renew=False, next_renewal=past, status="cancelled"))


def test_owner_ids_skip_garbage():
    good = "7b1e4b1c-6c2f-4c1e-9a7e-2f0c1d2e3f40"
    assert ps.owner_ids(_sub(owner_user_ids=f'["{good}", "nope"]')) == [good]
    assert ps.owner_ids(_sub(owner_user_ids="not json")) == []


def test_generated_payment_covers_one_cycle(fx):
    fx("")
    p = ps.payment_from_subscription(_sub(price=2346, cycle_months=12), date(2027, 1, 16),
                                     status="unconfirmed", created_by="system")
    assert (p.period_start, p.period_end) == (date(2027, 1, 16), date(2028, 1, 15))
    assert p.amount_aed == 2346 and p.status == "unconfirmed"


def test_card_last4_rejects_anything_but_four_digits():
    from app.schemas.schemas import PaymentCardIn
    assert PaymentCardIn(label="Corp Visa", last4="4821").last4 == "4821"
    assert PaymentCardIn(label="Corp Visa", last4="").last4 is None
    for bad in ("4111111111111111", "48a1", "482"):
        with pytest.raises(ValidationError):
            PaymentCardIn(label="Corp Visa", last4=bad)


def test_report_has_every_sheet_and_flags_leavers(fx):
    fx("")
    from app.services.purchases_report import build_report

    seat = SimpleNamespace(netbox_contact_id=42, person_name="Jane Doe", person_department="Newsroom",
                           login=None, person_email="jane@example.com", assigned_on=date(2026, 1, 5), notes=None)
    sub = _sub(seats=[seat], seats_total=12, next_renewal=date.today() + timedelta(days=10))
    pay = SimpleNamespace(paid_on=date(2026, 5, 3), vendor="Adobe", description="CC", category="software",
                          budget=None, payment_method="card", card=None, amount=100.0, currency="AED",
                          amount_aed=100.0, period_start=None, period_end=None, invoice_no=None,
                          status="unconfirmed", notes=None)
    wb = load_workbook(io.BytesIO(build_report([sub], [pay], {}, date(2026, 1, 1), date(2026, 12, 31), {42})))
    assert wb.sheetnames == ["Summary", "Payments", "Subscriptions", "Seats"]
    assert wb["Seats"]["G2"].value == "LEFT — revoke"
    assert wb["Subscriptions"]["W2"].value == "≤30d"
    assert wb["Payments"]["K2"].value == 100.0
    assert wb["Payments"]["K3"].value == "=SUM(K2:K2)"
