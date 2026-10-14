"""Excel report for Subscriptions & Purchases.

Same shape as the spreadsheet it replaces (payment log, monthly and vendor
summaries, license tracker) plus budgets and seats. Values are written already
computed — no TEXT()/EDATE() formulas, which broke on a Russian-locale Excel
("мммм") and silently drifted when someone typed over them.
"""
import io
from collections import defaultdict
from datetime import date
from typing import Iterable

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from app.models.models import PurchasePayment, Subscription
from app.services import purchases_service as ps

_HEADER_FONT = Font(bold=True, color="FFFFFF")
_HEADER_FILL = PatternFill("solid", fgColor="2F3B52")
_TOTAL_FONT = Font(bold=True)
_WARN_FILL = PatternFill("solid", fgColor="FDE7C7")
_BAD_FILL = PatternFill("solid", fgColor="F8D0CC")
_MONEY = "#,##0.00"
_DATE = "yyyy-mm-dd"

_METHOD = {"card": "Card", "bank_transfer": "Bank transfer", "paypal": "PayPal", "other": "Other"}
_CATEGORY = {"software": "Software", "cloud": "Cloud & hosting", "isp": "Internet (ISP)", "telecom": "Telecom & SIM",
             "security": "Security", "hardware": "Hardware", "services": "Services", "other": "Other"}


def _cycle(months: int) -> str:
    return {1: "Monthly", 3: "Quarterly", 6: "Half-yearly", 12: "Yearly", 24: "2 years", 36: "3 years"}.get(
        months, f"Every {months} months")


def _card(card) -> str:
    if not card:
        return ""
    return f"{card.label} ••{card.last4}" if card.last4 else card.label


def _sheet(wb: Workbook, title: str, headers: list[str], rows: Iterable[list], *,
           money_cols: Iterable[int] = (), date_cols: Iterable[int] = (), widths: dict[int, int] | None = None):
    ws = wb.create_sheet()
    ws.title = title
    ws.append(headers)
    for cell in ws[1]:
        cell.font, cell.fill = _HEADER_FONT, _HEADER_FILL
        cell.alignment = Alignment(vertical="center")
    for row in rows:
        ws.append(row)
    for col in money_cols:
        for (cell,) in ws.iter_rows(min_row=2, min_col=col, max_col=col):
            cell.number_format = _MONEY
    for col in date_cols:
        for (cell,) in ws.iter_rows(min_row=2, min_col=col, max_col=col):
            cell.number_format = _DATE
    for i, h in enumerate(headers, start=1):
        ws.column_dimensions[get_column_letter(i)].width = (widths or {}).get(i, max(12, min(len(h) + 4, 40)))
    ws.freeze_panes = "A2"
    if ws.max_row > 1:
        ws.auto_filter.ref = ws.dimensions
    return ws


def _total_row(ws, label_col: int, sum_cols: list[int]):
    last = ws.max_row
    row = ws.max_row + 1
    ws.cell(row=row, column=label_col, value="TOTAL").font = _TOTAL_FONT
    for col in sum_cols:
        letter = get_column_letter(col)
        c = ws.cell(row=row, column=col, value=f"=SUM({letter}2:{letter}{last})")
        c.font, c.number_format = _TOTAL_FONT, _MONEY


def build_report(subs: list[Subscription], payments: list[PurchasePayment], names: dict[str, str],
                 date_from: date, date_to: date, leaver_contact_ids: set[int]) -> bytes:
    today = date.today()
    wb = Workbook()

    # ── Summary ─────────────────────────────────────────
    total_aed = sum(p.amount_aed for p in payments)
    by_month: dict[date, list] = defaultdict(lambda: [0.0, 0])
    by_vendor: dict[str, list] = defaultdict(lambda: [0.0, 0])
    by_budget: dict[str, float] = defaultdict(float)
    by_category: dict[str, float] = defaultdict(float)
    for p in payments:
        m = p.paid_on.replace(day=1)
        by_month[m][0] += p.amount_aed
        by_month[m][1] += 1
        by_vendor[p.vendor][0] += p.amount_aed
        by_vendor[p.vendor][1] += 1
        by_budget[p.budget.name if p.budget else "—"] += p.amount_aed
        by_category[p.category] += p.amount_aed
    unconfirmed = sum(1 for p in payments if p.status == "unconfirmed")

    ws = wb.active
    ws.title = "Summary"
    ws.append(["IT purchases report"])
    ws["A1"].font = Font(bold=True, size=14)
    ws.append(["Period", f"{date_from:%d %b %Y} — {date_to:%d %b %Y}"])
    ws.append(["Generated", today.strftime("%d %b %Y")])
    ws.append(["Total spend, AED", round(total_aed, 2)])
    ws.append(["Total spend, USD", ps.aed_to_usd(total_aed)])
    ws.append(["Payments", len(payments)])
    if unconfirmed:
        ws.append(["Not yet confirmed", unconfirmed, "expected amounts — confirm against the bank statement"])
        ws.cell(row=ws.max_row, column=1).fill = _WARN_FILL
    active = [s for s in subs if s.status == "active"]
    run_rate = sum(m for s in active if (m := ps.monthly_aed(s)) is not None)
    ws.append(["Current run-rate / month, AED", round(run_rate, 2)])
    ws.append(["Current run-rate / month, USD", ps.aed_to_usd(run_rate)])
    ws.append(["Current run-rate / year, AED", round(run_rate * 12, 2)])
    ws.append(["Current run-rate / year, USD", ps.aed_to_usd(run_rate * 12)])
    for (cell,) in ws.iter_rows(min_row=4, max_row=ws.max_row, min_col=2, max_col=2):
        if isinstance(cell.value, float):
            cell.number_format = _MONEY
    ws.column_dimensions["A"].width = 34
    ws.column_dimensions["B"].width = 28

    def block(title: str, header: list[str], rows: list[list]):
        ws.append([])
        ws.append([title])
        ws.cell(row=ws.max_row, column=1).font = Font(bold=True, size=12)
        ws.append(header)
        for cell in ws[ws.max_row]:
            cell.font, cell.fill = _HEADER_FONT, _HEADER_FILL
        for r in rows:
            ws.append(r)
            for c in ws[ws.max_row][1:3]:
                if isinstance(c.value, float):
                    c.number_format = _MONEY

    block("By month", ["Month", "AED", "USD", "Payments"],
          [[m.strftime("%B %Y"), round(v[0], 2), ps.aed_to_usd(v[0]), v[1]] for m, v in sorted(by_month.items())])
    block("By vendor", ["Vendor", "AED", "USD", "Payments"],
          [[k, round(v[0], 2), ps.aed_to_usd(v[0]), v[1]] for k, v in sorted(by_vendor.items(), key=lambda kv: -kv[1][0])])
    block("By budget", ["Budget", "AED", "USD"],
          [[k, round(v, 2), ps.aed_to_usd(v)] for k, v in sorted(by_budget.items(), key=lambda kv: -kv[1])])
    block("By category", ["Category", "AED", "USD"],
          [[_CATEGORY.get(k, k), round(v, 2), ps.aed_to_usd(v)] for k, v in sorted(by_category.items(), key=lambda kv: -kv[1])])

    # ── Payments log ────────────────────────────────────
    ws = _sheet(
        wb, "Payments",
        ["Date", "Month", "Vendor", "Description", "Category", "Budget", "Method", "Card",
         "Amount", "Currency", "AED", "USD", "Period start", "Period end", "Invoice #", "Status", "Notes"],
        ([p.paid_on, p.paid_on.strftime("%B %Y"), p.vendor, p.description or "", _CATEGORY.get(p.category, p.category),
          p.budget.name if p.budget else "", _METHOD.get(p.payment_method, p.payment_method), _card(p.card),
          p.amount, p.currency, p.amount_aed, ps.aed_to_usd(p.amount_aed),
          p.period_start, p.period_end, p.invoice_no or "", p.status, p.notes or ""] for p in payments),
        money_cols=(9, 11, 12), date_cols=(1, 13, 14), widths={3: 22, 4: 36, 17: 40},
    )
    for row in ws.iter_rows(min_row=2):
        if row[15].value == "unconfirmed":
            for c in row:
                c.fill = _WARN_FILL
    if payments:
        _total_row(ws, 3, [11, 12])

    # ── Subscriptions & license tracker ─────────────────
    def renewal_state(s: Subscription) -> str:
        if s.status != "active":
            return "Cancelled"
        if not s.next_renewal:
            return ""
        days = (s.next_renewal - today).days
        if days < 0:
            return "Expired" if not s.auto_renew else "Overdue"
        return "≤30d" if days <= 30 else "Active"

    def sub_row(s: Subscription) -> list:
        price_aed = ps.to_aed(s.price or 0, s.currency)
        monthly = ps.monthly_aed(s)
        owners = ", ".join(names[i] for i in ps.owner_ids(s) if i in names)
        seats = f"{len(s.seats)}/{s.seats_total}" if s.seats_total else (str(len(s.seats)) if s.seats else "")
        return [
            s.vendor, s.product or "", _CATEGORY.get(s.category, s.category), s.seller or "", owners, s.budget.name if s.budget else "",
            s.account_email or "", _METHOD.get(s.payment_method, s.payment_method), _card(s.card),
            _cycle(s.cycle_months or 1), s.price or 0, s.currency, price_aed, ps.aed_to_usd(price_aed),
            monthly, ps.aed_to_usd(monthly), None if monthly is None else round(monthly * 12, 2),
            None if monthly is None else ps.aed_to_usd(monthly * 12), seats,
            s.next_renewal, (s.next_renewal - today).days if s.next_renewal else None,
            "Yes" if s.auto_renew else "No", renewal_state(s), s.reference or "", s.notes or "",
        ]

    ws = _sheet(
        wb, "Subscriptions",
        ["Vendor", "Product", "Category", "Seller", "Owner", "Budget", "Account email", "Method", "Card",
         "Cycle", "Price", "Currency", "Price AED", "Price USD", "Monthly AED", "Monthly USD",
         "Yearly AED", "Yearly USD", "Seats", "Next renewal", "Days left", "Auto-renew", "State",
         "Reference", "Notes"],
        (sub_row(s) for s in sorted(subs, key=lambda s: (s.status != "active", s.vendor.lower()))),
        money_cols=(11, 13, 14, 15, 16, 17, 18), date_cols=(20,), widths={2: 30, 7: 26, 25: 40},
    )
    for row in ws.iter_rows(min_row=2):
        state = row[22].value
        if state in ("Expired", "Overdue"):
            row[22].fill = _BAD_FILL
        elif state == "≤30d":
            row[22].fill = _WARN_FILL
    if subs:
        _total_row(ws, 1, [15, 16, 17, 18])

    # ── Seats ───────────────────────────────────────────
    seat_rows = []
    for s in subs:
        for seat in s.seats:
            leaver = seat.netbox_contact_id in leaver_contact_ids if seat.netbox_contact_id else None
            seat_rows.append([s.vendor, s.product or "", seat.person_name, seat.person_department or "",
                              seat.login or seat.person_email or "", seat.assigned_on,
                              "LEFT — revoke" if leaver else "", seat.notes or ""])
    ws = _sheet(wb, "Seats", ["Vendor", "Product", "Person", "Department", "Login / email", "Assigned",
                              "Status", "Notes"],
                sorted(seat_rows, key=lambda r: (r[0].lower(), r[2].lower())),
                date_cols=(6,), widths={2: 30, 3: 26, 5: 32, 8: 40})
    for row in ws.iter_rows(min_row=2):
        if row[6].value:
            for c in row:
                c.fill = _BAD_FILL

    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()
